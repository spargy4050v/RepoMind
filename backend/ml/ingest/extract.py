"""Normalize supported local uploads into the existing raw project-record schema.

Extraction intentionally stops before scoring. A reviewer must confirm the returned
fields, particularly OCR/fuzzy values, before ``inference.score_uploaded_scheme``
performs its established validation and feature engineering.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

import pandas as pd

MAX_FILE_BYTES = 10 * 1024 * 1024
SUPPORTED_SUFFIXES = {".csv", ".xlsx", ".xls", ".docx", ".pdf", ".png", ".jpg", ".jpeg"}
RAW_FIELDS = ("project_id", "mp_constituency", "state", "work_category", "contractor_id", "sanctioned_cost_inr", "regional_baseline_cost_inr", "recommended_date", "sanction_date", "start_date", "completion_certified_date", "fund_release_date", "planned_duration_days", "latitude", "longitude")

ALIASES = {
    "project_id": ("project id", "project_id", "scheme id", "scheme_id"),
    "mp_constituency": ("mp constituency", "constituency"),
    "state": ("state",), "work_category": ("work category", "category", "work type"),
    "contractor_id": ("contractor id", "contractor", "agency"),
    "sanctioned_cost_inr": ("sanctioned cost", "sanction amount", "sanctioned_cost_inr"),
    "regional_baseline_cost_inr": ("regional baseline cost", "baseline cost", "regional_baseline_cost_inr"),
    "recommended_date": ("recommended date",), "sanction_date": ("sanction date",),
    "start_date": ("start date",), "completion_certified_date": ("completion certified date", "completion date", "certified completion"),
    "fund_release_date": ("fund release date", "payment date", "release date"),
    "planned_duration_days": ("planned duration days", "planned duration", "duration days"),
    "latitude": ("latitude", "lat"), "longitude": ("longitude", "lon", "lng"),
}


@dataclass(frozen=True)
class ExtractionResult:
    """Unscored extracted values plus audit source and extraction certainty."""
    format: str
    record: dict[str, str]
    confidence: dict[str, Literal["high", "low"]]
    source: dict[str, object]


def _normalise(value: object) -> str:
    return re.sub(r"\s+", " ", str(value).replace("₹", "").replace(",", "")).strip()


def _key(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value).lower()).strip()


def _from_pairs(pairs: list[tuple[object, object]], *, low: bool) -> tuple[dict[str, str], dict[str, Literal["high", "low"]]]:
    record: dict[str, str] = {}
    confidence: dict[str, Literal["high", "low"]] = {}
    for raw_key, raw_value in pairs:
        key = _key(raw_key)
        for field, aliases in ALIASES.items():
            if field in record:
                continue
            matched = key in aliases
            fuzzy = not matched and any(alias in key or key in alias for alias in aliases)
            if matched or fuzzy:
                value = _normalise(raw_value)
                if value:
                    record[field] = value
                    confidence[field] = "low" if low or fuzzy else "high"
                break
    return record, confidence


def _from_text(text: str, *, low: bool) -> tuple[dict[str, str], dict[str, Literal["high", "low"]]]:
    pairs: list[tuple[str, str]] = []
    for line in text.splitlines():
        if ":" in line:
            key, value = line.split(":", 1)
            pairs.append((key, value))
    return _from_pairs(pairs, low=low)


def _tabular(data: bytes, suffix: str) -> tuple[dict[str, str], dict[str, Literal["high", "low"]], list[list[str]]]:
    frame = pd.read_csv(io.BytesIO(data), header=None) if suffix == ".csv" else pd.read_excel(io.BytesIO(data), header=None)
    frame = frame.dropna(axis=0, how="all").dropna(axis=1, how="all")
    if frame.empty:
        return {}, {}, []
    rows = [[_normalise(value) for value in row] for row in frame.fillna("").values.tolist()]
    # Header row followed by a record, or an explicit field/value form sheet.
    pairs = list(zip(rows[0], rows[1])) if len(rows) >= 2 and len(rows[0]) == len(rows[1]) else []
    if not pairs:
        pairs = [(row[0], row[1]) for row in rows if len(row) >= 2]
    record, confidence = _from_pairs(pairs, low=False)
    return record, confidence, rows[:50]


def extract_upload(filename: str, data: bytes) -> ExtractionResult:
    """Extract one supported local file without guessing missing project fields.

    Structured files use exact labels where possible. Document OCR values are
    always low confidence; documents with no labelled project fields fail rather
    than fabricating a record.
    """
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError("Unsupported file type. Use CSV, Excel, DOCX, PDF, PNG, JPG, or JPEG.")
    if not data or len(data) > MAX_FILE_BYTES:
        raise ValueError("Upload must be between 1 byte and 10 MB.")
    tables: list[list[str]] = []
    text = ""
    if suffix in {".csv", ".xlsx", ".xls"}:
        try:
            record, confidence, tables = _tabular(data, suffix)
        except Exception as error:
            raise ValueError(f"Could not read the spreadsheet: {error}") from error
    elif suffix == ".docx":
        try:
            from docx import Document
            document = Document(io.BytesIO(data))
            tables = [[cell.text for cell in row.cells] for table in document.tables for row in table.rows]
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            record, confidence = _from_pairs([(row[0], row[1]) for row in tables if len(row) >= 2], low=False)
            fallback, fallback_confidence = _from_text(text, low=False)
            record.update({key: value for key, value in fallback.items() if key not in record})
            confidence.update({key: value for key, value in fallback_confidence.items() if key not in confidence})
        except ImportError as error:
            raise ValueError("DOCX support requires python-docx. Install backend requirements.") from error
        except Exception as error:
            raise ValueError(f"Could not read the DOCX file: {error}") from error
    else:
        ocr = suffix in {".png", ".jpg", ".jpeg"}
        try:
            if suffix == ".pdf":
                from pypdf import PdfReader
                reader = PdfReader(io.BytesIO(data))
                if len(reader.pages) > 10:
                    raise ValueError("PDFs are limited to 10 pages.")
                text = "\n".join(page.extract_text() or "" for page in reader.pages)
                ocr = not text.strip()
                if ocr:
                    from pdf2image import convert_from_bytes
                    import pytesseract
                    text = "\n".join(pytesseract.image_to_string(page) for page in convert_from_bytes(data, first_page=1, last_page=min(10, len(reader.pages))))
            else:
                from PIL import Image
                import pytesseract
                text = pytesseract.image_to_string(Image.open(io.BytesIO(data)))
            record, confidence = _from_text(text, low=ocr)
        except ImportError as error:
            raise ValueError("PDF/image support requires pypdf, pytesseract, and local Tesseract. Install backend requirements.") from error
        except ValueError:
            raise
        except Exception as error:
            raise ValueError(f"Could not read the document: {error}") from error
    record = {key: value for key, value in record.items() if not key.startswith("_ground_truth_")}
    if not record:
        raise ValueError("Could not extract labelled project data from this file.")
    return ExtractionResult(suffix.removeprefix("."), record, confidence, {"text": text[:50000], "tables": tables})
